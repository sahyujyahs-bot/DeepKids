"""Builds app data (data.json) and the collated table (DeepKids_Board_Topics_Collated.xlsx)."""
import json, os, sys
HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
from sci_l1_data import T
from boards import RES, WEEKS
from catalog import *
from ncert_evs_content import C as EVS
from c3_daily import X as C3X
from c3_time_cur import H as C3H, CUR as C3CUR, JOKE_BOOK
from c3_deep import THINK as C3T, WS as C3W
from c3_prep import GUIDE as C3G, PREP as C3P, SUBJ_GUIDE, SUBJ_PREP
from sci_books import NCERT as SCI_NCERT, BOOK as CAM_SCI, NIOS_A, NIOS_B
import glob

# chapter packs: day-by-day plans, worksheets, experiments, activities, voices (app/content/SPEC.md)
PACKS = {}
from validate_pack import check as check_pack
for f in sorted(glob.glob(os.path.join(HERE, "..", "content", "*", "*.json"))):
    if check_pack(f): print("skipping pack that fails validation:", os.path.relpath(f, HERE)); continue
    p = json.load(open(f)); PACKS[p["key"]] = {k: v for k, v in p.items() if k not in ("key", "board", "subj", "level", "n")}

items = []  # every learnable unit: a board chapter/strand/theme, or a DeepKids topic
def add(**k): items.append(k)

for (subj, lvl), (book, chs) in NCERT.items():
    if subj == "sci": chs = [t for _, t, _ in SCI_NCERT[lvl]]  # titles as printed in the uploaded books
    for i, title in enumerate(chs, 1):
        c = ({**{k: v for k, v in EVS[(lvl, i)].items() if k != "dk"}, **({"x": {**C3X[i], "hours": C3H[i], "cur": C3CUR[i], "deep": C3T[i], "wsd": C3W[i], "guide": C3G[i], "prep": C3P[i]}} if lvl == 3 and i in C3X else {})} if subj == "sci" and (lvl, i) in EVS else None)
        if subj == "sci": c = {**(c or {}), "covers": SCI_NCERT[lvl][i - 1][2], "unit": SCI_NCERT[lvl][i - 1][0]}
        add(k=f"ncert|{subj}|{lvl}|{i}", board="ncert", subj=subj, level=lvl, book=f"{book} · Class {lvl}", n=str(i), title=title,
            dk=NCERT_DK.get((lvl, i), []) if subj == "sci" else [], c=c)
for subj, (book, strands) in CAMBRIDGE.items():
    for lvl in LEVELS:
        if subj == "sci":  # units from the uploaded Learner's Books
            bk, units = CAM_SCI[lvl]
            for i, (u, strand, tops) in enumerate(units, 1):
                add(k=f"cam|sci|{lvl}|{i}", board="cambridge", subj="sci", level=lvl, book=bk.replace("Cambridge Primary Science ", ""), n=str(i), title=u.split(" ", 1)[1], dk=[],
                    c={"covers": "; ".join(tops), "topics": tops, "strand": strand})
            continue
        for i, s in enumerate(strands, 1):
            add(k=f"cam|{subj}|{lvl}|{i}", board="cambridge", subj=subj, level=lvl, book=f"{book} · Stage {lvl}", n="Strand", title=s, dk=[])
for i, (n, t, cov) in enumerate(NIOS_A, 1):  # NIOS Level A is equivalent to Class 3
    add(k=f"nios|sci|3|{i}", board="nios", subj="sci", level=3, book="Environmental Studies (A 102) · Level A", n=n, title=t, dk=[], c={"covers": cov})
for i, (n, t, cov) in enumerate(NIOS_B, 1):  # NIOS Level B is equivalent to Class 5
    add(k=f"nios|sci|5|{i}", board="nios", subj="sci", level=5, book="Environmental Studies · Level B", n=n, title=t, dk=[], c={"covers": cov})
for it in items:  # attach chapter packs
    pk = it["k"].replace("cam|", "cambridge|", 1)
    if pk in PACKS: it["c"] = {**(it.get("c") or {}), "p": PACKS[pk]}

# science resource library (app/resources/SPEC.md): only files that pass the check
from validate_resources import check as check_res
RES_DIR = os.path.join(HERE, "..", "resources"); RESOURCES = []; _seen = {}
for f in sorted(glob.glob(os.path.join(RES_DIR, "*.json"))):
    if f.endswith("concepts.json"): continue
    if check_res(f, dict(_seen)): print("skipping resource file that fails validation:", os.path.basename(f)); continue
    for r in json.load(open(f)):
        u = r["url"].rstrip("/").lower()
        if u in _seen: continue
        _seen[u] = 1; RESOURCES.append({k: r[k] for k in ("title", "url", "source", "type", "ages", "access", "lang", "desc", "concepts") } | ({"how": r["how"]} if r.get("how") else {}) | ({"chapters": r["chapters"]} if r.get("chapters") else {}))
RCONCEPTS = [{k: c[k] for k in ("id", "concept", "chapters", "plus")} for c in json.load(open(os.path.join(RES_DIR, "concepts.json")))]

# Which classes each DeepKids topic is offered to. For now every EVS topic is open to Classes 1-5;
# narrow a topic by listing it here, e.g. "B1": [4, 5].
TOPIC_CLS = {}
topics = []
for t in T:
    topics.append({"id": t["id"], "q": t["q"], "name": t["topic"], "sub": [s.strip("• ").strip() for s in t["subtopics"].split("\n") if s.strip()],
                   "summary": t["summary"], "oneness": t["oneness"], "weeks": WEEKS[t["id"]], "full": t["id"] == "A1", "subj": "sci", "cls": TOPIC_CLS.get(t["id"], [1, 2, 3, 4, 5]),
                   "res": [{"type": a, "title": b, "url": c} for a, b, c in RES.get(t["id"], [])]})

data = {"resources": RESOURCES, "rconcepts": RCONCEPTS, "prepgen": {"guide": SUBJ_GUIDE, "prep": SUBJ_PREP}, "jokebook": JOKE_BOOK, "levels": {str(k): {"cls": v[0], "age": v[1]} for k, v in LEVELS.items()}, "subjects": SUBJECTS, "items": items, "topics": topics,
        "notes": {"ncert": {f"{s}|{l}": n for (s, l), n in NCERT_NOTES.items()},
                  "nios": {**{s: f"NIOS {n}: Level A and B books to be added." for s, n in NIOS_PENDING.items()},
                           "sci|1": "NIOS starts at Level A, which is equivalent to Class 3.", "sci|2": "NIOS starts at Level A, which is equivalent to Class 3.",
                           "sci|4": "NIOS has no level for Class 4: Level A is Class 3 and Level B is Class 5."},
                  "cambridge": {"hin": "Cambridge Primary has no Hindi course. Pair with NCERT Hindi or DeepKids language plans."}}}
json.dump(data, open(os.path.join(HERE, "data.json"), "w"), ensure_ascii=False, separators=(",", ":"))

# ---------- collated spreadsheet ----------
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
F = "Arial"; HDR = PatternFill("solid", fgColor="1F3A5F")
TNAME = {t["id"]: t["topic"] for t in T}
BNAME = {"ncert": "NCERT / CBSE", "cambridge": "Cambridge Primary", "nios": "NIOS OBE"}
wb = Workbook(); ws = wb.active; ws.title = "All boards"
cols = ["Board", "Subject", "DeepKids level", "Class", "Age", "NIOS level", "Book / framework", "Chapter no.", "Chapter / strand / theme", "DeepKids Science L1 topics linked", "Verify"]
ws.append(cols)
def row(it):
    c, a = LEVELS[it["level"]]
    return [BNAME[it["board"]], SUBJECTS[it["subj"]], it["level"], c, a, nios_level(it["level"]), it["book"], it["n"], it["title"],
            "; ".join(f"{d} {TNAME[d]}" for d in it["dk"]), "From the uploaded book" if it["subj"] == "sci" else "Titles from secondary listings" if it["board"] == "ncert" else ("Unit names from Learner's Books to add" if it["board"] == "cambridge" else "Lesson list to add")]
for it in sorted(items, key=lambda x: (x["board"], x["subj"], x["level"])): ws.append(row(it))
for (s, l), note in NCERT_NOTES.items():
    c, a = LEVELS[l]; ws.append(["NCERT / CBSE", SUBJECTS[s], l, c, a, nios_level(l), "(no book)", "", note, "Use DeepKids Science Plus", ""])
for s, n in NIOS_PENDING.items():
    ws.append(["NIOS OBE", SUBJECTS[s], "1–5", "1–5", "6–10", "A / B", n, "", "Lesson list to be added from NIOS study material", "", "Pending"])
ws2 = wb.create_sheet("DeepKids Plus · Science L1")
ws2.append(["Topic ID", "Big question", "Topic", "DeepKids level", "Class", "Age", "Subtopics", "NCERT EVS chapters linked (Cl 3–5)"])
rev = {}
for (l, n), ids in NCERT_DK.items():
    for d in ids: rev.setdefault(d, []).append(f"Cl {l} Ch {n} {NCERT[('sci', l)][1][n-1]}")
for t in T:
    ws2.append([t["id"], t["q"], t["topic"], 1, 1, 6, t["subtopics"].replace("• ", ""), "; ".join(rev.get(t["id"], [])) or "Beyond board syllabus"])
ws4 = wb.create_sheet("NCERT EVS · games & projects")
ws4.append(["Level", "Class", "Ch", "Chapter", "What it covers", "Game", "How to play", "Experiments", "Activities", "Project", "Project steps", "You end up with", "Extra reading", "Other NCERT books (same class)", "DeepKids Plus topics"])
for (l, n), c in sorted(EVS.items()):
    ws4.append([l, l, n, NCERT[("sci", l)][1][n-1], c["covers"], c["game"][0], c["game"][1], "\n".join(c["exp"]), "\n".join(c["act"]), c["proj"][0], c["proj"][1], c["proj"][2], "\n".join(c["read"]), "\n".join(c["links"]), "; ".join(f"{d} {TNAME[d]}" for d in c["dk"])])
ws3 = wb.create_sheet("Level map")
ws3.append(["DeepKids level", "NCERT class", "Cambridge stage", "NIOS OBE", "Age"])
for l, (c, a) in LEVELS.items(): ws3.append([l, c, c, f"Prepares for Level {nios_level(l)}" + (" (≈Class 3)" if l <= 3 else " (≈Class 5)"), a])
for w in (ws, ws2, ws3, ws4):
    for cell in w[1]: cell.font = Font(name=F, bold=True, color="FFFFFF"); cell.fill = HDR; cell.alignment = Alignment(wrap_text=True)
    for r in w.iter_rows(min_row=2):
        for cell in r: cell.font = Font(name=F, size=10); cell.alignment = Alignment(wrap_text=True, vertical="top")
    w.freeze_panes = "A2"; w.auto_filter.ref = w.dimensions
for i, wdt in enumerate([14, 14, 9, 7, 6, 8, 34, 9, 40, 44, 26], 1): ws.column_dimensions[get_column_letter(i)].width = wdt
for i, wdt in enumerate([8, 20, 32, 9, 7, 6, 70, 60], 1): ws2.column_dimensions[get_column_letter(i)].width = wdt
for i, wdt in enumerate([12, 12, 14, 34, 8], 1): ws3.column_dimensions[get_column_letter(i)].width = wdt
for i, wdt in enumerate([6, 6, 4, 22, 40, 18, 60, 50, 45, 20, 50, 26, 34, 34, 26], 1): ws4.column_dimensions[get_column_letter(i)].width = wdt
out = os.path.join(HERE, "..", "DeepKids_Board_Topics_Collated.xlsx"); wb.save(out)
print(len(items), "items;", len(topics), "DeepKids topics; xlsx:", os.path.abspath(out))
