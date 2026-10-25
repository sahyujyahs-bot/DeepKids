"""Build the DeepKids Combined curriculum map, Classes 1-5: Science, Maths, History, Philosophy.

Two fundamental units:
  subtopic-level           e.g. sci.forces.friction.L3   (DeepKids Combined)
  syllabus-class-chapter   e.g. NCERT Class 3 Ch 4, Cambridge Stage 3 topic 6.5   (board tracks)
Each subtopic-level gets: its topic-level, class, syllabus chapters, prerequisites, go-deep and go-wide
subtopic-levels (also across subjects), the maths it uses, and the founder's notebook items it teaches.
Every item of the founder's 'DeepKids HomeSchooling Track' tabs must be taught somewhere (checked).

Run: python3 app/build/curriculum/build.py
Writes app/content/curriculum/{sci,math,hist,phil}.json and app/DeepKids_Curriculum_L1-5.xlsx
"""
import json, os, re, sys, copy
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
HERE = os.path.dirname(os.path.abspath(__file__)); BUILD = os.path.dirname(HERE); APP = os.path.dirname(BUILD)
sys.path.insert(0, HERE)
import sci_map, sci_more, math_map, math_more, hist_map, phil_map
SUBJ_NAME = {"sci": "Science", "math": "Maths", "hist": "History", "phil": "Philosophy"}
problems = []

# ---------- board chapters (maths keys are kept under "M:" so NCERT EVS N4-10 and NCERT Maths N4-10 stay apart) ----------
CH = {}
wb = openpyxl.load_workbook(os.path.join(APP, 'DeepKids_Science_Board_Topics_L1-5.xlsx'), read_only=True)
for r in list(wb['All topics'].iter_rows(values_only=True))[1:]:
    if r[8] != 'From the book' or not r[5]: continue
    if r[3] == 'NCERT / CBSE':
        n = r[5].split('Ch ')[-1]; CH[f'N{r[0]}-{n}'] = dict(board='NCERT', cls=r[0], ref=f'Class {r[0]} Ch {n}', title=r[6], covers=r[7] or '', item=f'ncert|sci|{r[0]}|{n}')
    elif r[3] == 'Cambridge Primary':
        t = r[5].split(' · ')[-1]; u = t.split('.')[0]
        CH[f'C{r[0]}-{t}'] = dict(board='Cambridge', cls=r[0], ref=f'Stage {r[0]} {t}', title=r[6], covers=r[5].split(' · ')[0], item=f'cam|sci|{r[0]}|{u}')
    else:
        n = r[5].replace('Lesson ', ''); lvl = 'A' if r[0] == 3 else 'B'
        CH[f'O{r[0]}-{n}'] = dict(board='NIOS', cls=r[0], ref=f'Level {lvl} Lesson {n}', title=r[6], covers=r[7] or '', item=f'nios|sci|{r[0]}|{n}')
KEYS = {'sci': set(CH), 'math': set(), 'hist': set(), 'phil': set()}
wb = openpyxl.load_workbook(os.path.join(APP, 'DeepKids_Board_Topics_Collated.xlsx'), read_only=True)
for r in list(wb['All boards'].iter_rows(values_only=True))[1:]:
    if r[0] == 'NCERT / CBSE' and r[1] == 'Maths':
        k = f'M:N{r[3]}-{r[7]}'; CH[k] = dict(board='NCERT', cls=int(r[3]), ref=f'Class {r[3]} Ch {r[7]}', title=r[8], covers=r[6] or '', item=f'ncert|math|{r[3]}|{r[7]}'); KEYS['math'].add(k)
CAMN = {"Nc": "Counting and sequences", "Ni": "Integers and powers (calculation)", "Nm": "Money", "Np": "Place value, ordering and rounding",
        "Nf": "Fractions, decimals, percentages, ratio", "Gt": "Time", "Gg": "Geometrical reasoning, shapes and measurements",
        "Gp": "Position and transformation", "Ss": "Statistics", "Sp": "Probability", "TWM": "Thinking and working mathematically"}
for st in range(1, 6):
    for c, nm_ in CAMN.items():
        k = f'M:C{st}-{c}'; CH[k] = dict(board='Cambridge', cls=st, ref=f'Stage {st} {c}', title=nm_, covers='Cambridge Primary Maths 0096 framework', item=f'cam|math|{st}|{math_map.CAM_STRAND[c]}'); KEYS['math'].add(k)
ck = lambda subj, k: ('M:' if subj == 'math' else '') + k.rstrip('?')

# ---------- notebook items ----------
NB = json.load(open(os.path.join(APP, 'content', 'curriculum', 'notebook_items.json')))
NBI = {x['id']: x for v in NB.values() for x in v}
def expand(spec):
    out = []
    for part in [p.strip() for p in (spec or '').split(',') if p.strip()]:
        m = re.match(r'(\w+):(\d+)(?:-(\d+))?$', part)
        if not m: problems.append(f'bad notebook reference {part}'); continue
        for i in range(int(m.group(2)), int(m.group(3) or m.group(2)) + 1):
            k = f'{m.group(1)}:{i}'
            if k in NBI: out.append(k)
            else: problems.append(f'unknown notebook item {k}')
    return out

# ---------- merge maps ----------
def merge(base, more):
    T = copy.deepcopy(base)
    if more:
        idx = {t[0]: t for t in T}
        for tid, subs in getattr(more, 'SUBS', {}).items():
            if tid not in idx: problems.append(f'additions for unknown topic {tid}'); continue
            idx[tid][3].extend(subs)
        T += getattr(more, 'TOPICS', [])
        notes = getattr(more, 'NOTES', {})
        S = {f'{t[0]}/{s[0]}': s for t in T for s in t[3]}
        for ref, spec in notes.items():
            if ref not in S: problems.append(f'notes for unknown subtopic {ref}'); continue
            S[ref][4]['notes'] = ((S[ref][4].get('notes') or '') + ',' + spec).strip(',')
    return T
MAPS = {'sci': merge(sci_map.TOPICS, sci_more), 'math': merge(math_map.TOPICS, math_more), 'hist': merge(hist_map.TOPICS, None), 'phil': merge(phil_map.TOPICS, None)}
T_, S_ = {}, {}
for subj, topics in MAPS.items():
    T_[subj], S_[subj] = {}, {}
    for tid, tname, strand, subs in topics:
        if tid in T_[subj]: problems.append(f'{subj}: duplicate topic {tid}')
        T_[subj][tid] = dict(id=tid, name=tname, strand=strand, subs=[s[0] for s in subs])
        for sid, sname, levels, chaps, links in subs:
            ref = f'{tid}/{sid}'
            if ref in S_[subj]: problems.append(f'{subj}: duplicate subtopic {ref}')
            S_[subj][ref] = dict(topic=tid, sid=sid, name=sname, levels=levels, chaps=chaps, links=links)

def resolve(subj, ref):
    """'topic/sub' in this subject, or 'math:topic/sub' etc. in another"""
    m = re.match(r'(sci|math|hist|phil):(.+)$', ref)
    s2, r2 = (m.group(1), m.group(2)) if m else (subj, ref)
    if r2 not in S_[s2]: problems.append(f'{subj}: unknown link {ref}'); return None, None
    return s2, r2
def lv(s2, r2, n, mode):
    L = sorted(S_[s2][r2]['levels'])
    if mode == 'pre': c = [l for l in L if l <= n]; return c[-1] if c else None
    if mode == 'deep': c = [l for l in L if l >= n]; return c[0] if c else L[-1]
    return min(L, key=lambda l: (abs(l - n), -l))
sid_of = lambda subj, ref, l: f'{subj}.{ref.replace("/", ".")}.L{l}'

# ---------- subtopic-levels ----------
OUT, TL, PLACED, COVER = {}, {}, {}, {}
for subj in MAPS:
    T, S = T_[subj], S_[subj]; out, placed = [], {}
    for ref, s in S.items():
        L = sorted(s['levels']); notes = expand(s['links'].get('notes'))
        for k in notes: COVER.setdefault(k, []).append(sid_of(subj, ref, L[0]))
        for k in s['chaps']:
            kk = ck(subj, k)
            if kk not in CH: problems.append(f'{subj} {ref}: unknown chapter {k}')
            elif CH[kk]['cls'] not in s['levels']: problems.append(f'{subj} {ref}: chapter {k} is Class {CH[kk]["cls"]} but the subtopic has no L{CH[kk]["cls"]}')
        for i, n in enumerate(L):
            chaps, verify = [], []
            for k in s['chaps']:
                kk = ck(subj, k)
                if kk in CH and CH[kk]['cls'] == n:
                    chaps.append(kk); placed.setdefault(kk, []).append(sid_of(subj, ref, n))
                    if k.endswith('?'): verify.append(kk)
            pre = [sid_of(subj, ref, L[i - 1])] if i else []
            for p in s['links'].get('pre', []):
                s2, r2 = resolve(subj, p)
                if not s2: continue
                m = lv(s2, r2, n, 'pre')
                if m is None: problems.append(f'{subj} {ref} L{n}: prerequisite {p} is only taught from L{min(S_[s2][r2]["levels"])}')
                else: pre.append(sid_of(s2, r2, m))
            deep = [sid_of(subj, ref, L[i + 1])] if i + 1 < len(L) else []
            for d in s['links'].get('deep', []):
                s2, r2 = resolve(subj, d)
                if s2: deep.append(sid_of(s2, r2, lv(s2, r2, n, 'deep')))
            wide = [sid_of(subj, f'{s["topic"]}/{o}', n) for o in T[s['topic']]['subs'] if o != s['sid'] and n in S[f'{s["topic"]}/{o}']['levels']]
            for w in s['links'].get('wide', []):
                s2, r2 = resolve(subj, w)
                if s2: wide.append(sid_of(s2, r2, lv(s2, r2, n, 'near')))
            maths = []
            for m_ in s['links'].get('math', []):
                s2, r2 = resolve('math', m_.replace('math:', ''))
                if not s2: continue
                ml = lv(s2, r2, n, 'pre') or lv(s2, r2, n, 'deep')
                maths.append(dict(id=sid_of('math', r2, ml), chapters=[ck('math', k) for k in S_['math'][r2]['chaps'] if CH.get(ck('math', k), {}).get('cls') == ml]))
            out.append(dict(id=sid_of(subj, ref, n), subject=subj, topic=s['topic'], topic_name=T[s['topic']]['name'], strand=T[s['topic']]['strand'],
                            topic_level=f'{subj}.{s["topic"]}.L{n}', subtopic=s['sid'], name=s['name'], level=n, classes=[n], covers=s['levels'][n],
                            chapters=chaps, verify=verify, prerequisites=pre, go_deep=deep, go_wide=wide, maths=maths,
                            notebook=notes, source='boards' if s['chaps'] else ('notes' if notes or s['links'].get('src') == 'notes' else 'dk')))
    for k in sorted(KEYS[subj] - set(placed)): problems.append(f'{subj}: board chapter not mapped: {k} {CH[k]["ref"]} {CH[k]["title"]}')
    tl = {}
    for x in out:
        t = tl.setdefault(x['topic_level'], dict(id=x['topic_level'], subject=subj, topic=x['topic'], name=x['topic_name'], strand=x['strand'], level=x['level'], classes=[x['level']], subtopic_levels=[], chapters=[]))
        t['subtopic_levels'].append(x['id']); t['chapters'] += [c for c in x['chapters'] if c not in t['chapters']]
    OUT[subj], TL[subj], PLACED[subj] = out, sorted(tl.values(), key=lambda t: (t['level'], t['topic'])), placed
used_in = {}
for subj in ('sci', 'hist', 'phil'):
    for x in OUT[subj]:
        for m in x['maths']: used_in.setdefault(m['id'], []).append(x['id'])
for x in OUT['math']: x['used_in_other_subjects'] = used_in.get(x['id'], [])
ALL = [x for v in OUT.values() for x in v]
for x in ALL: x['chapter_refs'] = [f"{CH[k]['board']} {CH[k]['ref']} {CH[k]['title']}" for k in x['chapters']]
missing = [k for k in NBI if k not in COVER]
for k in missing: problems.append(f'notebook item not taught anywhere: {k} {NBI[k]["text"][:80]}')

# ---------- write data ----------
D = os.path.join(APP, 'content', 'curriculum'); os.makedirs(D, exist_ok=True)
plain = lambda k: k[2:] if isinstance(k, str) and k.startswith('M:') else k
def clean(o):
    if isinstance(o, dict): return {plain(a): clean(v) for a, v in o.items()}
    if isinstance(o, list): return [clean(v) for v in o]
    return plain(o)
for subj in MAPS:
    json.dump(clean(dict(subject=subj, name=SUBJ_NAME[subj], topic_levels=TL[subj], subtopic_levels=OUT[subj],
                         chapters={k: dict(CH[k], subtopic_levels=PLACED[subj].get(k, [])) for k in sorted(KEYS[subj])},
                         date_fixes=hist_map.DATE_FIXES if subj == 'hist' else [])),
              open(os.path.join(D, f'{subj}.json'), 'w'), ensure_ascii=False, indent=1)

# ---------- review spreadsheet ----------
X = openpyxl.Workbook(); H = Font(bold=True, color='FFFFFF'); HF = PatternFill('solid', fgColor='2B2340'); W = Alignment(wrap_text=True, vertical='top')
NAME = {x['id']: f"{x['name']} L{x['level']}" + ('' if x['subject'] == 'sci' else f" ({SUBJ_NAME[x['subject']]})") for x in ALL}
for x in OUT['sci']: NAME[x['id']] = f"{x['name']} L{x['level']}"
nm = lambda ids: '\n'.join(NAME.get(i, i) for i in ids)
def sheet(title, head, rows, widths, freeze='C2'):
    ws = X.create_sheet(title[:31]); ws.append(head)
    for c in ws[1]: c.font = H; c.fill = HF; c.alignment = W
    for r in rows: ws.append(r)
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = W
    ws.freeze_panes = freeze; ws.auto_filter.ref = ws.dimensions
X.remove(X.active)
rd = X.create_sheet('Read me')
lines = ["DeepKids Combined curriculum map, Classes 1-5: Science, Maths, History, Philosophy (draft for review)", "",
 "Two units: SUBTOPIC-LEVEL (DeepKids Combined; e.g. Friction L3) and SYLLABUS-CLASS-CHAPTER (board tracks; e.g. Cambridge Stage 3 topic 6.5).",
 "Level n is taught in Class n. A topic-level (e.g. Forces and movement L3) groups the subtopic-levels of that topic at that level.",
 "Each subtopic-level lists: what it covers at that level, board chapters, prerequisites, go-deep and go-wide (also across subjects), the maths it uses, and the items from the founder's HomeSchooling Track sheet it teaches.",
 "Nothing from the HomeSchooling Track sheet is left out for being 'too advanced': each item is placed at the level where it can be delivered well (see 'Notebook coverage').",
 "Source: 'boards' = taught by NCERT/Cambridge/NIOS; 'notes' = from the founder's sheets; 'dk' = DeepKids addition.",
 "Verify column: NCERT Maths (Joyful Mathematics / Maths Mela) chapter contents were mapped from titles and the books' sequence; check against the books when uploaded.",
 "Cambridge Maths is mapped by framework sub-strand per stage. NIOS Maths lessons are not listed yet.", ""]
for subj in MAPS:
    lines.append(f"{SUBJ_NAME[subj]}: {len(T_[subj])} topics, {len(S_[subj])} subtopics, {len(OUT[subj])} subtopic-levels, {len(TL[subj])} topic-levels"
                 + (f"; board chapters mapped {len(PLACED[subj])}/{len(KEYS[subj])}" if KEYS[subj] else ""))
lines += [f"HomeSchooling Track items taught: {len(NBI) - len(missing)}/{len(NBI)}.", f"Open checks: {len(problems)} (see the Checks sheet)."]
for l in lines: rd.append([l])
rd.column_dimensions['A'].width = 150
for subj in MAPS:
    lab = SUBJ_NAME[subj]
    sheet(f'{lab} topic-levels', ['Topic-level', 'Strand', 'Class', 'Subtopic-levels', 'Board chapters'],
          [[f"{t['name']} L{t['level']}", t['strand'], t['level'], nm(t['subtopic_levels']), '\n'.join(f"{CH[k]['board']} {CH[k]['ref']} {CH[k]['title']}" for k in t['chapters'])] for t in TL[subj]],
          [34, 18, 7, 50, 60])
    head = ['Subtopic-level', 'Topic-level', 'Strand', 'Class', 'What this level covers', 'Board chapters', 'Verify', 'Prerequisites', 'Go deep', 'Go wide',
            'Used in other subjects' if subj == 'math' else 'Maths it uses', 'From the HomeSchooling Track', 'Source', 'ID']
    rows = [[f"{x['name']} L{x['level']}", f"{x['topic_name']} L{x['level']}", x['strand'], x['level'], x['covers'], '\n'.join(x['chapter_refs']), ', '.join(map(plain, x['verify'])),
             nm(x['prerequisites']), nm(x['go_deep']), nm(x['go_wide']),
             nm(x['used_in_other_subjects']) if subj == 'math' else '\n'.join(NAME.get(m['id'], m['id']) + (f" ({', '.join(CH[c]['board'] + ' ' + CH[c]['ref'] for c in m['chapters'])})" if m['chapters'] else '') for m in x['maths']),
             '\n'.join(NBI[k]['text'][:90] for k in x['notebook']), x['source'], x['id']]
            for x in sorted(OUT[subj], key=lambda x: (x['level'], x['strand'], x['topic'], x['subtopic']))]
    sheet(f'{lab} subtopic-levels', head, rows, [30, 28, 14, 6, 60, 40, 9, 32, 32, 32, 36, 40, 8, 30])
sheet('Board chapters', ['Board', 'Subject', 'Class', 'Chapter', 'Title', 'What it covers', 'Mapped to subtopic-levels'],
      [[CH[k]['board'], SUBJ_NAME[s], CH[k]['cls'], CH[k]['ref'], CH[k]['title'], CH[k]['covers'], nm(PLACED[s].get(k, []))]
       for s in ('sci', 'math') for k in sorted(KEYS[s], key=lambda k: (CH[k]['board'], CH[k]['cls'], k))], [11, 9, 6, 22, 34, 50, 50])
TAB = {"math": "Maths", "bio": "Biology", "phy": "Physics", "chem": "Chemistry", "hist": "History", "phil": "Philosophy", "sub": "History of Our Sub Continent"}
sheet('Notebook coverage', ['Tab', 'Item', 'Taught in (first level of each subtopic)'],
      [[TAB[k.split(':')[0]], NBI[k]['text'], nm(COVER.get(k, [])) or 'NOT YET PLACED'] for k in NBI], [16, 70, 60], freeze='B2')
sheet('History date fixes', ['Corrected in the map'], [[f] for f in hist_map.DATE_FIXES], [120], freeze='A2')
sheet('Checks', ['Open check'], [[p] for p in problems] or [['None']], [150], freeze='A2')
XL = os.path.join(APP, 'DeepKids_Curriculum_L1-5.xlsx'); X.save(XL)
old = os.path.join(APP, 'DeepKids_Curriculum_Science_Maths_L1-5.xlsx')
if os.path.exists(old): os.remove(old)
for subj in MAPS:
    print(f"{SUBJ_NAME[subj]}: {len(T_[subj])} topics, {len(S_[subj])} subtopics, {len(OUT[subj])} subtopic-levels"
          + (f", chapters mapped {len(PLACED[subj])}/{len(KEYS[subj])}" if KEYS[subj] else ""))
print(f"notebook items taught: {len(NBI) - len(missing)}/{len(NBI)}; checks: {len(problems)}")
for p in problems[:80]: print('  ', p)
