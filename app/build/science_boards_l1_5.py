# All board Science / EVS topics for DeepKids Levels 1–5 (Classes 1–5, ages 6–10), one sheet.
# Run from repo root: python3 app/build/science_boards_l1_5.py
import openpyxl, sys, os
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'DeepKids_Science_Board_Topics_L1-5.xlsx')

from sci_books import NCERT, BOOK, NIOS_A, NIOS_B

CAM_ALL = [('Thinking and working scientifically', 'Every stage', 'Ask, predict, plan a fair test, observe, measure, record, conclude (grows each stage)'),
           ('Science in context', 'Every stage', 'Science in daily life; scientists past and present; caring for the environment')]

NIOS_ALL = ['Human existence and the body', 'Hygiene and health', 'Family and kinship', 'Respect for women and elders', 'Festivals',
            'Planets and the sky', 'Weather', 'Basic needs and diet', 'Transport', 'Communication', 'Rights: education, information, consumer rights']

ROWS = []  # level, class, age, board, book, unit/strand, topic, covers, confidence
for lv in range(1, 6):
    age = lv + 5
    if lv <= 2:
        ROWS.append([lv, lv, age, 'NCERT / CBSE', '(no EVS book)', '', 'No EVS book in Classes 1–2 (NCF 2023)',
                     'EVS ideas sit inside the English (Mridang), Hindi (Sarangi) and Maths (Joyful Mathematics) books', 'Confirmed'])
    else:
        for i, (unit, title, cov) in enumerate(NCERT[lv], 1):
            ROWS.append([lv, lv, age, 'NCERT / CBSE', f'Our Wondrous World · Class {lv}', f'{unit} · Ch {i}', title, cov, 'From the book'])
    title, units = BOOK[lv]
    for u, strands, topics in units:
        for i, t in enumerate(topics, 1):
            ROWS.append([lv, lv, age, 'Cambridge Primary', title, f'Unit {u} · {u.split()[0]}.{i}', t, strands, 'From the book'])
    for strand, sub, t in CAM_ALL:
        ROWS.append([lv, lv, age, 'Cambridge Primary', f'All stages', strand, t, '', 'Runs through every stage'])
    if lv == 3:
        for n, t, cov in NIOS_A:
            ROWS.append([lv, lv, age, 'NIOS OBE', 'Level A · Environmental Studies (A 102), ≈ Class 3', 'Lesson ' + n, t, cov, 'From the book'])
    elif lv < 3:
        ROWS.append([lv, lv, age, 'NIOS OBE', '(no level)', '', 'NIOS starts at Level A (≈ Class 3)', 'Level 1–2 work prepares for Level A', 'From the book'])
    elif lv == 5:
        for n, t, cov in NIOS_B:
            ROWS.append([lv, lv, age, 'NIOS OBE', 'Level B · Environmental Studies, ≈ Class 5', 'Lesson ' + n, t, cov, 'From the book'])
    else:
        ROWS.append([lv, lv, age, 'NIOS OBE', '(no level)', '', 'No NIOS level for Class 4', 'Level A ≈ Class 3, Level B ≈ Class 5', 'From the book'])

# ---------- workbook
wb = openpyxl.Workbook()
HF = PatternFill('solid', fgColor='53A373'); HFont = Font(bold=True, color='FFFFFF')
LF = {1: 'FFF6D6', 2: 'FFE9DA', 3: 'FDE4EE', 4: 'E6E8F6', 5: 'E3F3EC'}
thin = Side(style='thin', color='DDDDDD'); BD = Border(top=thin, bottom=thin, left=thin, right=thin)
WRAP = Alignment(wrap_text=True, vertical='top')

def style(ws, widths, level_col=None):
    for c in ws[1]: c.fill = HF; c.font = HFont; c.alignment = WRAP; c.border = BD
    for row in ws.iter_rows(min_row=2):
        lv = row[level_col].value if level_col is not None else None
        for c in row:
            c.alignment = WRAP; c.border = BD
            if lv in LF: c.fill = PatternFill('solid', fgColor=LF[lv])
    for i, w in enumerate(widths): ws.column_dimensions[get_column_letter(i + 1)].width = w
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions

ws = wb.active; ws.title = 'All topics'
ws.append(['Level', 'Class', 'Age', 'Board', 'Book / framework', 'Chapter / strand', 'Topic', 'What it covers', 'Check'])
for r in ROWS: ws.append(r)
style(ws, [7, 7, 6, 16, 30, 32, 48, 60, 18], 0)

# side by side: one block per level, boards as columns
ws = wb.create_sheet('Side by side')
ws.append(['Level · Class · Age', 'NCERT / CBSE', 'Cambridge Primary', 'NIOS OBE'])
for lv in range(1, 6):
    col = {b: [r[6] if b != 'Cambridge Primary' else f'{r[5].split(" · ")[-1]} {r[6]}' if r[8] == 'From the book' else r[6] for r in ROWS if r[0] == lv and r[3] == b]
           for b in ('NCERT / CBSE', 'Cambridge Primary', 'NIOS OBE')}
    if lv >= 3: col['NCERT / CBSE'] = [f'{i+1}. {t}' for i, t in enumerate(col['NCERT / CBSE'])]
    h = max(len(v) for v in col.values())
    for i in range(h):
        ws.append([f'Level {lv} · Class {lv} · Age {lv+5}' if i == 0 else lv] +
                  [col[b][i] if i < len(col[b]) else '' for b in ('NCERT / CBSE', 'Cambridge Primary', 'NIOS OBE')])
style(ws, [22, 40, 60, 40], None)
for row in ws.iter_rows(min_row=2):
    v = row[0].value; lv = v if isinstance(v, int) else int(str(v).split()[1])
    for c in row: c.fill = PatternFill('solid', fgColor=LF[lv])
    if isinstance(v, int): row[0].value = ''
    else: row[0].font = Font(bold=True)

ws = wb.create_sheet('Counts')
ws.append(['Level', 'Class', 'Age', 'NCERT chapters', 'Cambridge topics', 'NIOS lessons'])
for lv in range(1, 6):
    n = lambda b: sum(1 for r in ROWS if r[0] == lv and r[3] == b and r[8] == 'From the book' and r[4] != '(no level)')
    ws.append([lv, lv, lv + 5, n('NCERT / CBSE'), n('Cambridge Primary'), n('NIOS OBE')])
style(ws, [8, 8, 6, 16, 22, 14], 0)

ws = wb.create_sheet('Notes')
for line in [
 'Science / EVS topics of NCERT, Cambridge Primary and NIOS for DeepKids Levels 1–5 (Classes 1–5, ages about 6–10).',
 'NCERT: "Our Wondrous World" Classes 3–5 (NCF-SE 2023, Reprint 2026-27), from the uploaded chapter PDFs. "What it covers" summarises each unit\'s Note to the Teacher. Classes 1–2 have no EVS book.',
 'Cambridge: units and topics from the uploaded Learner\'s Books. Books 2, 4 and 5 are 2nd editions (current 0097 framework). Books 1 and 3 are 1st editions (older framework), so the 2nd editions may group topics differently.',
 'Book 5 prints the strand column of units 4–6 shifted; strands here are matched to the topics.',
 'Cambridge also runs two strands through every stage: Thinking and working scientifically, and Science in context.',
 'NIOS: Open Basic Education Level A (A 102) is "equivalent to Class 3" per the book, so its 12 lessons sit at Level 3. NIOS has nothing below Level A. Level B (≈ Class 5, 15 lessons) from the uploaded book.',
 'Still to upload if wanted: Cambridge Learner\'s Books 1 and 3 (2nd edition).',
]: ws.append([line])
ws.column_dimensions['A'].width = 140
for c in ws['A']: c.alignment = Alignment(wrap_text=True)
wb.move_sheet('Notes', offset=-3)
wb.save(OUT); print(len(ROWS), os.path.abspath(OUT))
