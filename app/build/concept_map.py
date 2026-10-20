# DeepKids Plus Science L1 topics mapped to board topics (Levels 1–5) by concept: one row per concept,
# one column per source. Reads the two sheets built earlier. Run from repo root: python3 app/build/concept_map.py
import openpyxl, os
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

APP = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..')
OUT = os.path.join(APP, 'DeepKids_Science_Concept_Map.xlsx')

wb0 = openpyxl.load_workbook(os.path.join(APP, 'DeepKids_Board_Topics_Collated.xlsx'), read_only=True)
DK = {r[0]: r[2] for r in list(wb0['DeepKids Plus · Science L1'].iter_rows(values_only=True))[1:] if r and r[0]}

# Board items keyed like N3-4 (NCERT Class 3 Ch 4), C2-1.3 (Cambridge Stage 2 topic 1.3), O3-5 (NIOS Level A lesson 5)
BI = {}
wb1 = openpyxl.load_workbook(os.path.join(APP, 'DeepKids_Science_Board_Topics_L1-5.xlsx'), read_only=True)
for r in list(wb1['All topics'].iter_rows(values_only=True))[1:]:
    if r[8] != 'From the book' or r[4] == '(no level)': continue
    if r[3] == 'NCERT / CBSE': k, lab = f'N{r[0]}-{r[5].split("Ch ")[-1]}', f'Ch {r[5].split("Ch ")[-1]} {r[6]}'
    elif r[3] == 'Cambridge Primary': t = r[5].split(' · ')[-1]; k, lab = f'C{r[0]}-{t}', f'{t} {r[6]}'
    else: n = r[5].replace('Lesson ', ''); k, lab = f'O{r[0]}-{n}', f'Lesson {n} {r[6]}'
    BI[k] = (r[0], lab)

# concept, DeepKids topics, board items, note
M = [
 ('Living and non-living; signs of life', ['A1'], ['C1-1.1', 'C3-3.1', 'C3-3.2', 'C3-3.3', 'O3-4'], 'Plus questions whether the line is real'),
 ('Growing up and having young', ['A1'], ['C1-1.3', 'C2-4.2'], 'Plus has this only inside A1 "signs of life"'),
 ('How life began', ['A2'], [], ''),
 ('Matter and particles: what everything is made of', ['A3'], ['C4-3.1'], ''),
 ('Changing materials: heating, melting, reactions', ['A3'], ['C2-3.5', 'C4-3.2', 'C4-3.3', 'C4-3.4'], ''),
 ('Materials and their properties', ['A4'], ['N3-10', 'C1-4.1', 'C1-4.2', 'C1-4.3', 'C2-3.1', 'C2-3.2', 'C2-3.3', 'C2-3.4', 'C3-5.1', 'C3-5.2', 'C3-5.3', 'C3-5.4'], ''),
 ('Making things: houses, paper, cloth', ['A4'], ['N3-11', 'N4-8', 'N5-8'], 'Plus has materials but not making things from them'),
 ('Rocks and soil', ['A4'], ['C2-1.4', 'O5-8'], ''),
 ('Gases, evaporation, condensation, solutions', ['A4', 'B5'], ['C5-3.1', 'C5-3.2', 'C5-3.3', 'C5-3.4'], 'Plus touches this via water as a material and clouds/rain'),
 ('Plants: parts and what they need', ['A5'], ['N3-4', 'C1-2.1', 'C1-2.3', 'C3-1.1', 'C3-1.2', 'C3-1.3'], ''),
 ('Seeds, plant growth and plant life cycles', ['A5'], ['C1-2.2', 'C3-1.4', 'C5-1.1', 'C5-1.2', 'C5-1.3', 'C5-1.4'], ''),
 ('Kinds of living things and sorting them', ['A5'], ['N4-3', 'C2-4.1', 'C3-3.5'], ''),
 ('Habitats and adaptation', ['A5', 'B4'], ['N3-5', 'N4-9', 'N5-6', 'C1-1.2', 'C2-1.1', 'C2-1.2', 'C2-1.3', 'C4-4.4', 'C5-6.3', 'C5-6.4'], ''),
 ('Living in harmony: food chains and the web of life', ['C4', 'C12'], ['N3-6', 'N4-4', 'C4-2.4', 'O5-7'], 'Plus has no food-chain topic yet; closest is "what we share with other life"'),
 ('The universe, Sun, Moon, stars and planets', ['B1'], ['N4-10', 'C4-5.4', 'O3-6', 'O5-11'], ''),
 ('Gravity', ['B1', 'B6'], ['C5-5.1', 'C5-5.2'], ''),
 ('Planet Earth: inside, land, water and rivers', ['B1'], ['N5-2', 'C4-4.1', 'O5-5'], ''),
 ('Day, night, seasons and rhythms of nature', ['B5'], ['N5-9', 'C2-5.3', 'C4-5.5', 'C5-6.1', 'C5-6.2'], 'NCERT Class 4 Ch 10 and NIOS Lesson 6 also cover this (listed under the sky row)'),
 ('Water: sources, water cycle and saving it', ['B5', 'B8'], ['N3-7', 'N5-1'], ''),
 ('Nature\'s events: volcanoes, earthquakes, floods', ['B8'], ['C4-4.2', 'C4-4.3'], ''),
 ('Society, community, culture and festivals', ['B2'], ['N3-2', 'N3-3', 'N4-1', 'N4-2', 'N5-5', 'O3-2', 'O3-3', 'O3-11', 'O3-12', 'O5-1', 'O5-2', 'O5-3', 'O5-4', 'O5-6'], 'Boards treat this as social studies; Plus asks why societies are this way'),
 ('Transport and communication', ['B2'], ['O3-10'], 'NCERT Class 4 Ch 2 also covers this (listed under society)'),
 ('Microorganisms', ['B3', 'C1'], ['N5-3'], 'NCERT: microbes that spoil food, and good microbes'),
 ('The world of birds', ['B4'], [], 'Boards meet birds only inside habitats and nature trails'),
 ('Forces and movement', ['B6'], ['N4-7', 'C1-5.1', 'C1-5.2', 'C1-5.3', 'C1-5.4', 'C2-2.1', 'C2-2.2', 'C2-2.3', 'C2-2.4', 'C3-6.1', 'C3-6.2', 'C3-6.3', 'C3-6.4', 'C3-6.5', 'C5-5.3', 'C5-5.4'], 'NCERT: spinning, floating and sinking'),
 ('Magnets', ['B6'], ['C3-5.5', 'C5-5.5', 'C5-5.6'], 'Plus mentions electromagnetic force; no magnets topic yet'),
 ('Sound', ['B6'], ['C1-6.1', 'C1-6.3', 'C5-2.1', 'C5-2.2', 'C5-2.3', 'C5-2.4'], ''),
 ('Light, seeing and shadows', ['B7'], ['C2-5.1', 'C2-5.2', 'C4-5.1', 'C4-5.2', 'C4-5.3', 'C4-5.6'], 'NCERT Class 4 Ch 10 also has shadow lengths (listed under the sky row)'),
 ('Caring for the environment', ['B9'], ['N5-4', 'N5-10', 'C2-1.5', 'O3-5', 'O5-9'], ''),
 ('Waste: reduce, reuse, recycle', ['B10'], ['N3-12', 'O5-15'], 'NCERT Class 4 Ch 8 also teaches the 5Rs (listed under making things)'),
 ('Genes and inheriting features', ['C2'], ['C2-4.3'], ''),
 ('Where humans came from', ['C3'], [], ''),
 ('Me and others: alike and different', ['C4', 'C5'], ['C1-3.1', 'C1-3.2', 'C3-3.4'], ''),
 ('Our closest cousins', ['C5'], [], ''),
 ('Family', ['C6'], ['N3-1', 'O3-1'], ''),
 ('Body parts and skeleton', ['C7'], ['C1-3.3', 'C4-1.1', 'C4-1.2', 'C4-1.3', 'C4-1.4', 'O5-13'], ''),
 ('Senses', ['C8'], ['C1-3.4', 'C1-6.2', 'C3-4.1', 'C3-4.2', 'C3-4.3'], ''),
 ('Thinking and memory', ['C9'], [], ''),
 ('Health, hygiene, exercise and illness', ['C10'], ['N3-9', 'N4-6', 'C2-4.4', 'C2-4.5', 'C3-2.4', 'C4-1.5', 'O3-9', 'O5-14'], ''),
 ('Electricity and staying safe with it', ['C10'], ['C2-6.1', 'C2-6.2', 'C2-6.3', 'C4-6.1', 'C4-6.2', 'C4-6.3', 'C4-6.4', 'C4-6.5'], 'Plus covers only safety; boards go on to circuits'),
 ('Food and a balanced diet', ['C11'], ['N3-8', 'N4-5', 'C1-1.4', 'C3-2.1', 'C3-2.2', 'C3-2.3', 'C5-4.2', 'O3-8', 'O5-12'], ''),
 ('Digestion: how the body uses food', ['C11'], ['C5-4.1'], ''),
 ('Basic needs: air, water, food, clothing, shelter', ['C11', 'C12'], ['O3-7'], ''),
 ('Energy', ['C12'], ['N5-7', 'C4-2.1', 'C4-2.2', 'C4-2.3', 'O5-10'], ''),
]

used = [b for _, _, bs, _ in M for b in bs]
missing = sorted(set(BI) - set(used)); dup = sorted({b for b in used if used.count(b) > 1}); bad = sorted(set(used) - set(BI))
assert not missing and not dup and not bad, (missing, dup, bad)
assert set(DK) <= {d for _, ds, _, _ in M for d in ds}

def cell(bs, p):
    return '\n'.join(f'L{BI[b][0]} · {BI[b][1]}' for b in bs if b[0] == p) or '–'

def strength(bs):
    if not bs: return 'Plus only'
    lv = sorted({BI[b][0] for b in bs})
    return 'Level 1' if 1 in lv else 'Later level'

ROWS = []
for c, ds, bs, note in M:
    lv = sorted({BI[b][0] for b in bs})
    ROWS.append([c, '\n'.join(f'{d} {DK[d]}' for d in ds), cell(bs, 'N'), cell(bs, 'C'), cell(bs, 'O'),
                 ', '.join(f'L{x}' for x in lv) or '–', strength(bs), note])

wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Concept map'
ws.append(['Concept', 'DeepKids Plus (Level 1)', 'NCERT / CBSE', 'Cambridge Primary', 'NIOS OBE', 'Board levels', 'Board match', 'Note'])
for r in ROWS: ws.append(r)
HF = PatternFill('solid', fgColor='53A373'); thin = Side(style='thin', color='DDDDDD'); BD = Border(top=thin, bottom=thin, left=thin, right=thin)
FILL = {'Level 1': 'E3F3EC', 'Later level': 'FFF6D6', 'Plus only': 'FDE4EE'}
for c in ws[1]: c.fill = HF; c.font = Font(bold=True, color='FFFFFF'); c.alignment = Alignment(wrap_text=True, vertical='top'); c.border = BD
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical='top'); c.border = BD
    row[0].font = Font(bold=True); row[6].fill = PatternFill('solid', fgColor=FILL[row[6].value])
for i, w in enumerate([30, 34, 34, 44, 34, 12, 12, 36]): ws.column_dimensions[get_column_letter(i + 1)].width = w
ws.freeze_panes = 'B2'; ws.auto_filter.ref = ws.dimensions

ws = wb.create_sheet('Notes')
cnt = {k: sum(1 for r in ROWS if r[6] == k) for k in FILL}
for line in [
 'Each row is one concept. The four columns show where that concept sits in DeepKids Plus and in each board; "L3" means Level 3 (Class 3, about age 8).',
 'Every board topic from the Board Topics L1-5 sheet appears exactly once. Every DeepKids Plus Level 1 topic appears at least once.',
 'Where a DeepKids topic has no Level 1 board match, it is mapped to the level where the boards teach that concept.',
 f'Board match: Level 1 = a board teaches it at Level 1 ({cnt["Level 1"]} concepts) · Later level = boards teach it only at Levels 2–5 ({cnt["Later level"]}) · Plus only = no board teaches it ({cnt["Plus only"]}).',
 'Sources: NCERT Our Wondrous World 3–5 (2026-27), Cambridge Learner\'s Books 1–5 (Books 1 and 3 are 1st editions), NIOS Environmental Studies Levels A and B.',
]: ws.append([line])
ws.column_dimensions['A'].width = 140
for c in ws['A']: c.alignment = Alignment(wrap_text=True)
wb.save(OUT); print(len(ROWS), cnt)
